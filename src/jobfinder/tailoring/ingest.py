from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument

from jobfinder import paths
from jobfinder.llm.base import LLMProvider, load_prompt
from jobfinder.tailoring.master_schema import MasterResume


@dataclass
class IngestResult:
    resume_path: Path
    cover_path: Path | None
    warnings: list[str] = field(default_factory=list)


def extract_docx_text(path: Path) -> str:
    doc = DocxDocument(str(path))
    lines: list[str] = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def find_input_files(input_dir: Path) -> tuple[Path | None, Path | None]:
    docx = sorted(p for p in input_dir.glob("*.docx") if not p.name.startswith("~$"))
    cover = next((p for p in docx if "cover" in p.name.lower() or "lettre" in p.name.lower()), None)
    rest = [p for p in docx if p != cover]
    resume = next((p for p in rest if "resume" in p.name.lower()), None)
    if resume is None:
        resume = next((p for p in rest if "cv" in p.name.lower()), None)
    if resume is None and rest:
        resume = rest[0]
    return resume, cover


def _resolve_explicit(path: Path | None, input_dir: Path) -> Path | None:
    if path is None:
        return None
    if not path.is_absolute() and not path.exists():
        path = input_dir / path
    if not path.exists():
        raise FileNotFoundError(f"no such file: {path}")
    return path


def run_ingest(
    *,
    llm: LLMProvider,
    input_dir: Path | None = None,
    master_dir: Path | None = None,
    resume_file: Path | None = None,
    cover_file: Path | None = None,
) -> IngestResult:
    input_dir = input_dir or paths.input_dir()
    master_dir = master_dir or paths.master_dir()
    resume_file = _resolve_explicit(resume_file, input_dir)
    cover_file = _resolve_explicit(cover_file, input_dir)
    if resume_file is None or cover_file is None:
        found_resume, found_cover = find_input_files(input_dir)
        resume_file = resume_file or found_resume
        cover_file = cover_file or found_cover
    if resume_file is None:
        raise FileNotFoundError(f"no resume .docx found in {input_dir}")
    warnings: list[str] = []
    text = extract_docx_text(resume_file)
    data = llm.complete_json(
        task="ingest_resume",
        system=load_prompt("ingest_resume"),
        user=text,
        schema=MasterResume.model_json_schema(),
        tier="strong",
    )
    master = MasterResume.model_validate(data)
    if not master.experience:
        warnings.append("no experience entries extracted — check the docx formatting")
    master_dir.mkdir(parents=True, exist_ok=True)
    resume_out = master_dir / "resume.yaml"
    master.to_yaml(resume_out)
    cover_out: Path | None = None
    if cover_file is not None:
        cover_out = master_dir / "cover-letter.md"
        cover_out.write_text(extract_docx_text(cover_file) + "\n", encoding="utf-8")
    else:
        warnings.append("no cover letter .docx found (name must contain 'cover' or 'lettre')")
    return IngestResult(resume_out, cover_out, warnings)
