from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document as DocxDocument

from jobfinder.tailoring.render_docx import HEADINGS, docx_text
from jobfinder.tailoring.render_pdf import pdf_pages
from jobfinder.tailoring.requirements import coverage_ratio, keyword_coverage

ATS_READY_THRESHOLD = 75
ALLOWED_FONTS = {
    "Calibri", "Arial", "Helvetica", "Times New Roman", "Georgia", "Garamond", "Cambria", None,
}
REQUIRED_HEADINGS = ("summary", "experience", "education")


def inspect_docx(path: Path) -> dict:
    doc = DocxDocument(str(path))
    with zipfile.ZipFile(str(path)) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
    header_footer = any(
        p.text.strip()
        for s in doc.sections
        for p in list(s.header.paragraphs) + list(s.footer.paragraphs)
    )
    bad_fonts = sorted(
        {r.font.name for p in doc.paragraphs for r in p.runs if r.font.name not in ALLOWED_FONTS}
    )
    if doc.styles["Normal"].font.name not in ALLOWED_FONTS:
        bad_fonts.append(str(doc.styles["Normal"].font.name))
    texts = [p.text.strip() for p in doc.paragraphs]
    headings_found = [t for t in texts if t and t == t.upper() and len(t) <= 20]
    return {
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes) + xml.count("<w:drawing"),
        "text_boxes": xml.count("<w:txbxContent"),
        "columns": xml.count('w:num="2"') + xml.count('w:num="3"'),
        "header_footer_text": header_footer,
        "bad_fonts": bad_fonts,
        "headings_found": headings_found,
        "words": len(docx_text(path).split()),
    }


def ats_report(
    docx_path: Path, pdf_path: Path | None, required_terms: list[str], lang: str = "en"
) -> dict:
    info = inspect_docx(docx_path)
    violations = [k for k in ("tables", "images", "text_boxes", "columns") if info[k]]
    if info["header_footer_text"]:
        violations.append("header_footer_text")
    if info["bad_fonts"]:
        violations.append("fonts")
    expected = {HEADINGS.get(lang, HEADINGS["en"])[k] for k in REQUIRED_HEADINGS}
    if not expected <= set(info["headings_found"]):
        violations.append("missing_headings")
    matched, missing = keyword_coverage(docx_text(docx_path), required_terms)
    coverage = coverage_ratio(matched, missing)
    pages = pdf_pages(pdf_path) if pdf_path else max(1, round(info["words"] / 550 + 0.49))
    score = 100 - 15 * len(violations) - round((1 - coverage) * 40) - (10 if pages > 2 else 0)
    return {
        **info,
        "violations": violations,
        "keyword_matched": matched,
        "keyword_missing": missing,
        "coverage": round(coverage, 3),
        "pages": pages,
        "score": max(0, min(100, score)),
    }
