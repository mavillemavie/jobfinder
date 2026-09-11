from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from jobfinder.tailoring.master_schema import MasterResume

HEADINGS: dict[str, dict[str, str]] = {
    "en": {
        "summary": "SUMMARY", "skills": "SKILLS", "experience": "EXPERIENCE",
        "education": "EDUCATION", "certifications": "CERTIFICATIONS", "languages": "LANGUAGES",
        "projects": "PROJECTS",
    },
    "fr": {
        "summary": "SOMMAIRE", "skills": "COMPÉTENCES", "experience": "EXPÉRIENCE",
        "education": "FORMATION", "certifications": "CERTIFICATIONS", "languages": "LANGUES",
        "projects": "PROJETS",
    },
}
BODY_FONT = "Calibri"
BODY_PT = 10.5


def _new_document() -> DocxDocument:
    doc = DocxDocument()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_PT)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(54)
        section.left_margin = section.right_margin = Pt(58)
    return doc


def _heading(doc: DocxDocument, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)


def _contact_line(resume: MasterResume) -> str:
    c = resume.contact
    return " | ".join(x for x in [c.location, c.phone, c.email, c.linkedin, c.website] if x)


def _name_block(doc: DocxDocument, resume: MasterResume) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(resume.contact.name)
    run.bold = True
    run.font.size = Pt(16)
    doc.add_paragraph(_contact_line(resume))


def _dates(start: str | None, end: str | None) -> str:
    return f"{start or ''} – {end or ''}".strip(" –")


def render_resume_docx(resume: MasterResume, path: Path, lang: str = "en") -> Path:
    h = HEADINGS.get(lang, HEADINGS["en"])
    doc = _new_document()
    _name_block(doc, resume)
    if resume.summary:
        _heading(doc, h["summary"])
        doc.add_paragraph(resume.summary.strip())
    if resume.skills:
        _heading(doc, h["skills"])
        for g in resume.skills:
            p = doc.add_paragraph()
            r = p.add_run(f"{g.category}: ")
            r.bold = True
            p.add_run(", ".join(g.items))
    if resume.experience:
        _heading(doc, h["experience"])
        for e in resume.experience:
            p = doc.add_paragraph()
            r = p.add_run(f"{e.title} — {e.company}")
            r.bold = True
            meta = " | ".join(x for x in [e.location, _dates(e.start, e.end)] if x)
            if meta:
                p.add_run(f"  ({meta})")
            for b in e.bullets:
                doc.add_paragraph(b, style="List Bullet")
    if resume.education:
        _heading(doc, h["education"])
        for ed in resume.education:
            bits = [ed.credential, ed.field, ed.institution, ed.year]
            doc.add_paragraph(", ".join(str(b) for b in bits if b))
    if resume.certifications:
        _heading(doc, h["certifications"])
        for c in resume.certifications:
            doc.add_paragraph(c, style="List Bullet")
    if resume.languages:
        _heading(doc, h["languages"])
        doc.add_paragraph(
            ", ".join(f"{lg.name} ({lg.level})" if lg.level else lg.name for lg in resume.languages)
        )
    if resume.projects:
        _heading(doc, h["projects"])
        for pr in resume.projects:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(f"{pr.name}: ")
            r.bold = True
            p.add_run(pr.description or "")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def render_cover_letter_docx(text: str, resume: MasterResume, path: Path) -> Path:
    doc = _new_document()
    _name_block(doc, resume)
    doc.add_paragraph("")
    for para in [p.strip() for p in text.split("\n\n") if p.strip()]:
        doc.add_paragraph(para)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def docx_text(path: Path) -> str:
    doc = DocxDocument(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
