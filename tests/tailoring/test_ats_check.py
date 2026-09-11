from pathlib import Path

from docx import Document as DocxDocument

from jobfinder.tailoring.ats_check import ATS_READY_THRESHOLD, ats_report, inspect_docx
from jobfinder.tailoring.master_schema import MasterResume
from jobfinder.tailoring.render_docx import render_resume_docx
from jobfinder.tailoring.render_pdf import render_resume_pdf

FIXTURE = Path(__file__).parent.parent / "fixtures" / "master" / "resume.yaml"


def test_clean_resume_scores_high(tmp_path: Path) -> None:
    m = MasterResume.from_yaml(FIXTURE)
    docx = render_resume_docx(m, tmp_path / "r.docx")
    pdf = render_resume_pdf(m, tmp_path / "r.pdf")
    rep = ats_report(docx, pdf, ["sql", "power bi", "python", "looker"])
    assert rep["violations"] == [] and rep["pages"] == 1
    assert rep["keyword_missing"] == ["looker"] and rep["coverage"] == 0.75
    assert rep["score"] >= ATS_READY_THRESHOLD


def test_table_and_header_are_violations(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("Name")
    doc.add_table(rows=1, cols=2)
    doc.sections[0].header.paragraphs[0].text = "phone 555"
    p = tmp_path / "bad.docx"
    doc.save(str(p))
    info = inspect_docx(p)
    assert info["tables"] == 1 and info["header_footer_text"] is True
    rep = ats_report(p, None, ["sql"])
    assert {"tables", "header_footer_text", "missing_headings"} <= set(rep["violations"])
    assert rep["score"] < ATS_READY_THRESHOLD
