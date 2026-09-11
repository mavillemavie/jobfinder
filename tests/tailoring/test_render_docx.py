from pathlib import Path

from docx import Document as DocxDocument

from jobfinder.tailoring.master_schema import MasterResume
from jobfinder.tailoring.render_docx import docx_text, render_cover_letter_docx, render_resume_docx

FIXTURE = Path(__file__).parent.parent / "fixtures" / "master" / "resume.yaml"


def test_resume_docx_structure(tmp_path: Path) -> None:
    m = MasterResume.from_yaml(FIXTURE)
    out = render_resume_docx(m, tmp_path / "r.docx")
    doc = DocxDocument(str(out))
    assert doc.tables == [] and len(doc.inline_shapes) == 0
    paras = [p for p in doc.paragraphs if p.text.strip()]
    assert paras[0].text == "Jordan Test" and paras[0].runs[0].bold
    assert paras[0].runs[0].font.size.pt == 16
    texts = [p.text for p in paras]
    assert "SUMMARY" in texts and "EXPERIENCE" in texts and "EDUCATION" in texts
    assert "SKILLS" in texts
    bullets = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert any("Power BI dashboards" in p.text for p in bullets)
    assert all(
        p.text == "" for s in doc.sections for p in s.header.paragraphs + s.footer.paragraphs
    )
    assert doc.styles["Normal"].font.name == "Calibri"
    assert doc.styles["Normal"].font.size.pt == 10.5
    assert "jordan@example.com" in docx_text(out)


def test_french_headings(tmp_path: Path) -> None:
    m = MasterResume.from_yaml(FIXTURE)
    texts = docx_text(render_resume_docx(m, tmp_path / "fr.docx", lang="fr"))
    assert "EXPÉRIENCE" in texts and "FORMATION" in texts and "COMPÉTENCES" in texts


def test_cover_letter_docx(tmp_path: Path) -> None:
    m = MasterResume.from_yaml(FIXTURE)
    out = render_cover_letter_docx(
        "Dear Hiring Manager,\n\nFirst paragraph.\n\nSecond paragraph.", m, tmp_path / "c.docx"
    )
    text = docx_text(out)
    assert text.startswith("Jordan Test") and "Second paragraph." in text
