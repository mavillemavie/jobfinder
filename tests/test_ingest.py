from pathlib import Path

from docx import Document as DocxDocument
from typer.testing import CliRunner

from jobfinder import paths
from jobfinder.cli import app
from jobfinder.llm.fake import FakeLLM
from jobfinder.tailoring.ingest import extract_docx_text, find_input_files, run_ingest
from jobfinder.tailoring.master_schema import MasterResume

FIXTURE = Path(__file__).parent / "fixtures" / "master" / "resume.yaml"


def _make_docx(path: Path, lines: list[str]) -> None:
    doc = DocxDocument()
    for line in lines:
        doc.add_paragraph(line)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "SQL"
    table.rows[0].cells[1].text = "Power BI"
    doc.save(path)


def test_extract_docx_text_includes_tables(tmp_path: Path) -> None:
    p = tmp_path / "r.docx"
    _make_docx(p, ["Jordan Test", "Reporting Analyst"])
    text = extract_docx_text(p)
    assert "Jordan Test" in text and "Power BI" in text


def test_find_input_files_by_name(tmp_path: Path) -> None:
    (tmp_path / "JF Resume 2026.docx").write_bytes(b"")
    (tmp_path / "Cover Letter.docx").write_bytes(b"")
    resume, cover = find_input_files(tmp_path)
    assert resume.name.startswith("JF Resume") and cover.name == "Cover Letter.docx"


def test_find_input_files_prefers_resume_over_cv(tmp_path: Path) -> None:
    (tmp_path / "A_CV.docx").write_bytes(b"")
    (tmp_path / "JF Resume.docx").write_bytes(b"")
    resume, cover = find_input_files(tmp_path)
    assert resume.name == "JF Resume.docx"


def test_run_ingest_writes_master(home: Path) -> None:
    _make_docx(paths.input_dir() / "resume.docx", ["Jordan Test", "Reporting Analyst"])
    _make_docx(paths.input_dir() / "cover letter.docx", ["Dear Hiring Manager,", "I am writing..."])
    canned = MasterResume.from_yaml(FIXTURE).model_dump(mode="json")
    llm = FakeLLM({"ingest_resume": canned})
    result = run_ingest(llm=llm)
    assert result.resume_path.exists() and result.cover_path.exists()
    assert "Dear Hiring Manager" in result.cover_path.read_text()
    assert llm.calls[0]["tier"] == "strong"
    assert "Reporting Analyst" in llm.calls[0]["user"]


def test_run_ingest_explicit_paths_override_heuristic(home: Path) -> None:
    _make_docx(paths.input_dir() / "resume.docx", ["Jordan Test", "Reporting Analyst"])
    chosen_resume = paths.input_dir() / "not-obviously-named.docx"
    _make_docx(chosen_resume, ["Explicit Chosen Person", "Data Wizard"])
    _make_docx(paths.input_dir() / "cover letter.docx", ["Dear Hiring Manager,", "I am writing..."])
    chosen_cover = paths.input_dir() / "another-letter.docx"
    _make_docx(chosen_cover, ["Dear Explicit Cover,", "Chosen letter body"])
    canned = MasterResume.from_yaml(FIXTURE).model_dump(mode="json")
    llm = FakeLLM({"ingest_resume": canned})
    result = run_ingest(llm=llm, resume_file=chosen_resume, cover_file=chosen_cover)
    assert "Explicit Chosen Person" in llm.calls[0]["user"]
    assert "Jordan Test" not in llm.calls[0]["user"]
    assert "Dear Explicit Cover" in result.cover_path.read_text()


def test_cli_ingest_without_input_exits_1(home: Path) -> None:
    result = CliRunner().invoke(app, ["ingest", "--provider", "fake"])
    assert result.exit_code == 1
    assert "input/" in result.stdout


def test_explicit_relative_path_resolves_against_input_dir(home: Path) -> None:
    _make_docx(paths.input_dir() / "r.docx", ["Resume Holder", "Senior Developer"])
    _make_docx(paths.input_dir() / "c.docx", ["Dear Manager,", "I am excited..."])
    canned = MasterResume.from_yaml(FIXTURE).model_dump(mode="json")
    llm = FakeLLM({"ingest_resume": canned})
    result = run_ingest(llm=llm, resume_file=Path("r.docx"), cover_file=Path("c.docx"))
    assert result.resume_path.exists() and result.cover_path.exists()
    assert "Resume Holder" in llm.calls[0]["user"]
    assert "Senior Developer" in llm.calls[0]["user"]
    assert "Dear Manager" in result.cover_path.read_text()


def test_explicit_missing_path_raises_with_path(home: Path) -> None:
    llm = FakeLLM({})
    try:
        run_ingest(llm=llm, resume_file=Path("nope.docx"))
        raise AssertionError("expected FileNotFoundError")
    except FileNotFoundError as exc:
        assert "nope.docx" in str(exc)


def test_master_dir_override_is_honoured(home: Path, tmp_path: Path) -> None:
    _make_docx(paths.input_dir() / "resume.docx", ["Alternate Resume", "Data Engineer"])
    _make_docx(paths.input_dir() / "cover letter.docx", ["Dear Team,", "Application body"])
    canned = MasterResume.from_yaml(FIXTURE).model_dump(mode="json")
    llm = FakeLLM({"ingest_resume": canned})
    elsewhere = tmp_path / "elsewhere"
    run_ingest(llm=llm, master_dir=elsewhere)
    assert (elsewhere / "resume.yaml").exists()
    assert (elsewhere / "cover-letter.md").exists()
    assert not (paths.master_dir() / "resume.yaml").exists()
